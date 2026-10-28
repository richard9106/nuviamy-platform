"""Synthetic data-return rehearsal: no production records or external services."""
import hashlib
import json
import time
from datetime import timedelta
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.export import build_export_data, build_practice_export, ExportFileUnavailable
from apps.accounts.models import UserProfile
from apps.audit.models import AuditLog
from apps.billing.models import PracticeSubscription
from apps.clients.models import Client
from apps.dashboard.models import SupportTicket, SupportMessage
from apps.documents.models import ClientDocument
from apps.portal.models import (ClientPortalAccess, ClientPortalRequest, PortalConversation,
    PortalMessage, IntakePacketTemplate, ClientIntakeAssignment)
from apps.practices.models import Practice, TherapistProfile, ExternalIntegration


@override_settings(MFA_REQUIRED=False, MFA_PRIVILEGED_REQUIRED=False,
                   EMAIL_TWO_STEP_REQUIRED=False, STRIPE_SECRET_KEY='',
                   EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend',
                   FIELD_ENCRYPTION_KEY='MDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDA=')
class DataReturnTests(TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory(prefix='nuviamy-return-synthetic-')
        self.addCleanup(self.directory.cleanup)
        self.storage = override_settings(MEDIA_ROOT=self.directory.name,
            STORAGES={'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
                      'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}})
        self.storage.enable()
        self.addCleanup(self.storage.disable)
        self.practice = Practice.objects.create(name='Synthetic return practice')
        self.user = get_user_model().objects.create_user('return-owner', 'owner@example.com', 'SyntheticPass123!')
        UserProfile.objects.create(user=self.user, practice=self.practice, role=UserProfile.Role.OWNER)
        self.therapist = TherapistProfile.objects.create(user=self.user,practice=self.practice,
            license_number='SYNTHETIC-RETURN',license_state='CA')
        self.patient = Client.objects.create(practice=self.practice,primary_therapist=self.therapist,
            first_name='Synthetic',last_name='Return',emergency_contact_name='Synthetic contact',
            emergency_contact_phone='555-0100',address_line1='Synthetic address')
        self.document = ClientDocument.objects.create(practice=self.practice,client=self.patient,
            title='Synthetic consent',description='Synthetic document description',
            original_filename='consent.txt',file=ContentFile(b'Synthetic consent bytes',name='consent.txt'))
        self.patient.profile_photo.save('synthetic-avatar.png',ContentFile(b'Synthetic avatar bytes'))
        self.template = IntakePacketTemplate.objects.create(practice=self.practice,name='Synthetic template',
            questions=['Synthetic question'])
        self.assignment = ClientIntakeAssignment.objects.create(practice=self.practice,client=self.patient,
            template=self.template,answers={'question_0':'Synthetic answer'},status='submitted')
        self.conversation = PortalConversation.objects.create(practice=self.practice,client=self.patient,
            subject='Synthetic conversation')
        PortalMessage.objects.create(practice=self.practice,conversation=self.conversation,author=self.user,
            author_kind='practice',body='Synthetic patient message')
        ClientPortalRequest.objects.create(practice=self.practice,client=self.patient,category='general',
            subject='Synthetic request',message='Synthetic request content')
        self.client.force_login(self.user)

    def archive(self, **kwargs):
        return self.client.post(reverse('practice_data_export'),{'scope':'all',**kwargs})

    def test_return_archive_captures_portal_forms_contacts_and_file_integrity(self):
        response = self.archive()
        self.assertEqual(response.status_code,200)
        self.assertIn('no-store',response['Cache-Control'])
        with ZipFile(BytesIO(response.content)) as archive:
            records = archive.read('practice-data.json')
            data = json.loads(records)
            manifest = json.loads(archive.read('manifest.json'))
            self.assertEqual(manifest['schema_version'],2)
            self.assertEqual(manifest['records']['sha256'],hashlib.sha256(records).hexdigest())
            self.assertEqual(data['intake_assignments'][0]['answers'],{'question_0':'Synthetic answer'})
            self.assertEqual(data['intake_assignments'][0]['template_snapshot']['name'],'Synthetic template')
            self.assertEqual(data['portal_messages'][0]['body'],'Synthetic patient message')
            self.assertEqual(data['portal_requests'][0]['message'],'Synthetic request content')
            self.assertEqual(data['clients'][0]['emergency_contact_name'],'Synthetic contact')
            self.assertEqual(data['clients'][0]['address_line1'],'Synthetic address')
            self.assertEqual(data['documents'][0]['description'],'Synthetic document description')
            self.assertEqual(data['therapists'][0]['license_number'],'SYNTHETIC-RETURN')
            self.assertEqual(data['team_members'][0]['id'],self.user.pk)
            self.assertIn('README.txt',archive.namelist())
            self.assertEqual(len(manifest['files']),2)
            for entry in manifest['files']:
                content=archive.read(entry['path'])
                self.assertEqual(entry['bytes'],len(content))
                self.assertEqual(entry['sha256'],hashlib.sha256(content).hexdigest())
                self.assertNotIn('..',entry['path'])
                self.assertNotIn('consent.txt',entry['path'])
            for section,count in manifest['record_counts'].items():
                self.assertEqual(count,len(data[section]))
        self.assertTrue(AuditLog.objects.filter(action=AuditLog.Action.EXPORT).exists())
        self.assertTrue(Client.objects.filter(pk=self.patient.pk).exists())
        self.assertTrue(self.document.file.storage.exists(self.document.file.name))

    def test_other_tenant_and_authentication_secrets_never_enter_archive(self):
        other=Practice.objects.create(name='Other synthetic clinic')
        Client.objects.create(practice=other,first_name='FOREIGN',last_name='PRIVATE')
        ClientIntakeAssignment.objects.create(practice=other,
            client=Client.objects.get(practice=other),template=IntakePacketTemplate.objects.create(
                practice=other,name='FOREIGN TEMPLATE',questions=['FOREIGN QUESTION']),
            answers={'question_0':'FOREIGN ANSWER'})
        ExternalIntegration.objects.create(practice=self.practice,provider='google',
            access_token='NEVER-EXPORT-OAUTH',refresh_token='NEVER-EXPORT-REFRESH')
        ticket=SupportTicket.objects.create(practice=self.practice,requester=self.user,
            contact_email=self.user.email,topic='technical')
        SupportMessage.objects.create(ticket=ticket,body='PRIVATE-ACCOUNT-SUPPORT')
        response=self.archive()
        with ZipFile(BytesIO(response.content)) as archive:
            raw=archive.read('practice-data.json').decode()
        for secret in ('FOREIGN',self.user.password,'NEVER-EXPORT-OAUTH','NEVER-EXPORT-REFRESH',
                       'PRIVATE-ACCOUNT-SUPPORT','must_change_password','access_token','refresh_token'):
            self.assertNotIn(secret,raw)

    def test_missing_blob_prevents_partial_zip_and_success_audit(self):
        Path(self.document.file.path).unlink()
        response=self.archive()
        self.assertRedirects(response,reverse('profile_settings'))
        self.assertNotEqual(response.get('Content-Type'),'application/zip')
        self.assertFalse(AuditLog.objects.filter(action=AuditLog.Action.EXPORT).exists())
        self.assertTrue(AuditLog.objects.filter(metadata__event='export_file_unavailable').exists())
        self.assertTrue(ClientDocument.objects.filter(pk=self.document.pk).exists())
        self.assertTrue(Client.objects.filter(pk=self.patient.pk).exists())
        self.assertNotIn(self.document.file.name,response.content.decode())

    def test_storage_failure_and_empty_reference_are_not_silently_skipped(self):
        with patch.object(ClientDocument._meta.get_field('file').storage,'open',side_effect=OSError('private path')):
            with self.assertRaises(ExportFileUnavailable):
                build_practice_export(self.practice,'all')
        self.document.file=''
        self.document.save(update_fields=['file'])
        with self.assertRaises(ExportFileUnavailable):
            build_practice_export(self.practice,'all')

    def test_billing_scope_excludes_clinical_forms_messages_and_files(self):
        response=self.archive(scope='billing')
        with ZipFile(BytesIO(response.content)) as archive:
            data=json.loads(archive.read('practice-data.json'))
            self.assertNotIn('clients',data)
            self.assertNotIn('intake_assignments',data)
            self.assertNotIn('portal_messages',data)
            self.assertEqual(json.loads(archive.read('manifest.json'))['files'],[])

    def test_excel_text_is_not_an_executable_formula(self):
        self.patient.first_name='=HYPERLINK("https://example.invalid","synthetic")'
        self.patient.save(update_fields=['first_name'])
        response=self.archive(format='xlsx')
        workbook=load_workbook(BytesIO(response.content),data_only=False)
        cells=list(workbook['clients'].iter_rows())
        column=[cell.value for cell in cells[0]].index('first_name')
        self.assertEqual(cells[1][column].data_type,'s')
        self.assertTrue(cells[1][column].value.startswith("'="))
        self.assertEqual(build_export_data(self.practice,'all')['clients'][0]['first_name'],self.patient.first_name)

    def test_expiry_and_cancellation_keep_read_export_and_support_without_erasure(self):
        self.practice.subscription_required=True
        self.practice.free_trial_ends_at=timezone.now()-timedelta(days=1)
        self.practice.save(update_fields=['subscription_required','free_trial_ends_at'])
        ticket=SupportTicket.objects.create(practice=self.practice,requester=self.user,
            contact_email=self.user.email,topic='other')
        for stage in ('trial_expired','subscription_canceled'):
            with self.subTest(stage=stage):
                self.assertEqual(self.client.get(reverse('clients:list')).status_code,200)
                self.assertEqual(self.archive().status_code,200)
                self.assertRedirects(self.client.post(reverse('clients:delete',args=[self.patient.pk])),
                    reverse('billing:subscription_access'))
                self.assertTrue(Client.objects.filter(pk=self.patient.pk).exists())
                self.assertTrue(ClientDocument.objects.filter(pk=self.document.pk).exists())
                response=self.client.post(reverse('support_request_detail',args=[ticket.pk]),
                    {'message':'Synthetic closure review question'})
                self.assertRedirects(response,reverse('support_request_detail',args=[ticket.pk]))
            PracticeSubscription.objects.get_or_create(practice=self.practice,defaults={
                'status':'canceled','stripe_subscription_id':'sub_synthetic_canceled',
                'cancellation_access_until':timezone.now()-timedelta(days=1)})

    def test_therapist_and_patient_cannot_download_practice_return(self):
        profile=self.user.nuvia_profile
        profile.role=UserProfile.Role.THERAPIST
        profile.save(update_fields=['role'])
        self.assertEqual(self.archive().status_code,403)
        patient_user=get_user_model().objects.create_user('synthetic-return-patient')
        ClientPortalAccess.objects.create(practice=self.practice,client=self.patient,user=patient_user)
        self.client.force_login(patient_user)
        self.assertEqual(self.archive().status_code,403)

    def test_export_requires_fresh_password_even_without_mandatory_authenticator(self):
        session=self.client.session
        session['security_auth_at']=time.time()-settings.SECURITY_REAUTH_TIMEOUT-1
        session.save()
        response=self.archive()
        self.assertEqual(response.status_code,302)
        self.assertTrue(response.url.startswith(reverse('security_reauthenticate')))
        self.assertIn('practice-data-export',response.url)
        # Password-only users are not forced into TOTP enrollment by this safeguard.
        prompt=self.client.get(response.url)
        self.assertContains(prompt,'name="password"')
        self.assertNotContains(prompt,'name="code"')
        confirmed=self.client.post(response.url,{'password':'SyntheticPass123!'})
        self.assertRedirects(confirmed,reverse('profile_settings')+'#practice-data-export')
        self.assertEqual(self.archive().status_code,200)
